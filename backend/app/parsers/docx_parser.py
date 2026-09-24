"""DOCX question extraction via python-docx. Splits paragraphs into question
blocks (PARSER_QUESTION_PATTERNS), extracts options per block
(PARSER_OPTION_PATTERNS), and detects the correct option via, in order:
a trailing "Javoblar:" answer key in the paragraph text, an answer-key
*table* (a real format seen in the wild: a "Savol"/"Javob" row pair with
question numbers and letters), or bold-run fallback. Never auto-clears
needs_review here — that decision belongs to the teacher review step (see
docs/spec.md section 4).
"""

import io
import re

from docx import Document

from app.parsers.base import BaseParser, ParsedOption, ParsedQuestion
from app.parsers.text_split import (
    extract_answer_key,
    question_number,
    split_block,
    split_into_blocks,
    split_off_answer_key,
)

_ANSWER_ROW_LABELS = {"javob", "javoblar"}
_QUESTION_ROW_LABELS = {"savol", "savollar"}


def _extract_table_answer_key(document: Document) -> dict[int, str]:
    """Some exam files carry the answer key as a table instead of inline
    text: one row of question numbers headed "Savol", immediately followed
    by a row of option letters headed "Javob" (repeated in blocks of N
    columns for a wide exam). python-docx only reads paragraph text by
    default, so tables need this separate pass."""
    key: dict[int, str] = {}
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for i in range(len(rows) - 1):
            header, values = rows[i], rows[i + 1]
            if not header or not values:
                continue
            if header[0].strip().lower() not in _QUESTION_ROW_LABELS:
                continue
            if values[0].strip().lower() not in _ANSWER_ROW_LABELS:
                continue
            for num_text, letter in zip(header[1:], values[1:]):
                if num_text.strip().isdigit() and re.match(r"^[A-Da-d]$", letter.strip()):
                    key[int(num_text.strip())] = letter.strip().upper()
    return key


class DocxParser(BaseParser):
    def parse(self, file_bytes: bytes) -> list[ParsedQuestion]:
        document = Document(io.BytesIO(file_bytes))

        lines: list[dict] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            bold = any(run.bold for run in paragraph.runs if run.text.strip())
            lines.append({"text": text, "bold": bool(bold)})

        body_lines, key_lines = split_off_answer_key(lines)
        answer_key = {**_extract_table_answer_key(document), **extract_answer_key(key_lines)}
        blocks = split_into_blocks(body_lines)

        questions: list[ParsedQuestion] = []
        for block in blocks:
            prompt, option_lines = split_block(block)
            if not prompt or not option_lines:
                continue

            num = question_number(block)
            correct_letter = answer_key.get(num) if num is not None else None

            options: list[ParsedOption] = []
            for i, opt in enumerate(option_lines):
                letter = chr(ord("A") + i)
                is_correct = (correct_letter == letter) if correct_letter else opt["bold"]
                options.append(ParsedOption(text=opt["text"], is_correct=is_correct))

            confidence = "high" if (correct_letter and len(options) == 4) else "low"
            questions.append(ParsedQuestion(prompt=prompt, options=options, confidence=confidence))

        return questions
